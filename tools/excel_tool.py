"""
tools/excel_tool.py - Excel report generation tool for Agent01.

Uses openpyxl to create formatted Excel workbooks from CSV datasets or records
and save them into the output directory.
"""

import os
from typing import Optional, List, Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from tools.dataset_tool import read_dataset


def generate_excel_report(
    source_file: str = "data/sample_employees.csv",
    output_path: str = "output/employee_report.xlsx"
) -> Dict[str, Any]:
    """
    Generates an Excel spreadsheet (.xlsx) from an employee CSV dataset.

    Args:
        source_file: Path to the input CSV dataset.
        output_path: Target path for the generated Excel file.

    Returns:
        Dictionary containing execution summary, row count, and output path.

    Raises:
        FileNotFoundError: If the source CSV file does not exist.
        ValueError: If the dataset is empty.
    """
    # 1. Read records from source dataset
    records = read_dataset(file_path=source_file)
    if not records:
        raise ValueError(f"Cannot generate Excel report: source dataset '{source_file}' has no records.")

    # 2. Ensure output directory exists
    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    # 3. Create openpyxl workbook and active sheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"

    # 4. Write header row
    headers = list(records[0].keys())
    ws.append(headers)

    # Style header row (Dark blue background, bold white text)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # 5. Write data rows
    for row in records:
        row_values = []
        for h in headers:
            val = row.get(h, "")
            # Convert numeric strings (like salary) to numbers if possible
            if isinstance(val, str) and val.isdigit():
                val = int(val)
            row_values.append(val)
        ws.append(row_values)

    # 6. Adjust column widths automatically
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[col_letter].width = max(max_length + 3, 12)

    # 7. Save the workbook
    try:
        wb.save(output_path)
    except Exception as err:
        raise IOError(f"Failed to save Excel file to '{output_path}': {err}")

    return {
        "success": True,
        "records_written": len(records),
        "output_path": output_path,
        "message": f"Successfully created Excel report at '{output_path}' with {len(records)} records."
    }
